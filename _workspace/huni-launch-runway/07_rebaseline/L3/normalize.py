#!/usr/bin/env python3
# L3 — 기존 자산 3종 → 단일 스키마 정규화 (판정·병합 없음)
# 원판정은 원천 값을 그대로 옮기고, 원본 레코드 전체를 raw 에 보존한다.
import csv, json, os

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '../../../..'))
OUT = os.path.join(ROOT, '_workspace/huni-launch-runway/07_rebaseline/L3')

P_LEDGER = '_workspace/huni-launch-runway/01_scope/ledger.json'
P_CANON  = '_workspace/huni-launch-scope/01_foundation/ia-feature-canon.csv'
P_FITGAP = '_workspace/huni-launch-scope/02_gap/fit-gap-matrix.csv'
P_XLSX   = 'docs/huni/후니프린팅_통합IA_일정_역할분담_260616.xlsx'

def rp(p): return os.path.join(ROOT, p)

rows = []

# ── 원천 1: runway ledger 410건 ──────────────────────────────────
# legacy_id = 원본 id 그대로 (대시보드 localStorage 체크상태 보존 → 재발급 금지)
ledger = json.load(open(rp(P_LEDGER), encoding='utf-8'))
for r in ledger:
    rows.append({
        'legacy_id': r['id'],
        '원천': 'runway',
        '원본ID': r['id'],
        '기능서술': r.get('title', ''),
        '원판정': r.get('status_now', ''),
        '원판정상세': {
            'status_now': r.get('status_now'),
            'verdict': r.get('verdict'),
            'kind': r.get('kind'),
            'raw_kind': r.get('raw_kind'),
            'phase': r.get('phase'),
            'effort_hint': r.get('effort_hint'),
            'origin': r.get('origin'),
            'conflict': r.get('conflict', False),
            'conflict_note': r.get('conflict_note'),
        },
        '근거': [s.get('evidence') for s in r.get('sources', []) if s.get('evidence')],
        '돈여부': r.get('money_critical'),
        '주문여부': r.get('order_critical'),
        '의존': {
            'blocked_by': r.get('blocked_by', []),
            'external_dep': r.get('external_dep', []),
            'precondition': r.get('precondition'),
        },
        'raw': r,
    })

# ── 원천 2: launch-scope 162건 (canon + fit-gap, no 기준 1:1) ─────
canon = {x['no']: x for x in csv.DictReader(open(rp(P_CANON), encoding='utf-8'))}
fitgap = {x['no']: x for x in csv.DictReader(open(rp(P_FITGAP), encoding='utf-8'))}
assert set(canon) == set(fitgap), 'canon/fit-gap no 집합 불일치'
for i, no in enumerate(sorted(canon, key=int), start=1):
    c, g = canon[no], fitgap[no]
    line = int(no) + 1  # csv 헤더 1행
    rows.append({
        'legacy_id': f'SCOPE-{int(no):03d}',
        '원천': 'scope',
        '원본ID': f'launch-scope#no={no}',
        '기능서술': ' / '.join(x for x in (c.get('시스템'), c.get('영역'), c.get('기능')) if x),
        '원판정': g.get('판정', ''),
        '원판정상세': {
            '판정': g.get('판정'),
            '우선순위': c.get('우선순위'),
            'phase_canon': c.get('phase'),
            'phase_fitgap': g.get('phase'),
            '담당': c.get('담당'),
            '규모': c.get('규모'),
            '구분': c.get('구분'),
            'launch_flag': c.get('launch_flag'),
            '개발위치': g.get('개발위치'),
            '개발방안요약': g.get('개발방안요약'),
        },
        '근거': [f'{P_CANON}:{line}', f'{P_FITGAP}:{line}',
                 g.get('shopby근거') or None],
        '돈여부': None,   # 원천에 플래그 없음 — 판정 금지(L4의 일)
        '주문여부': None,
        '의존': {
            'blocked_by': [],
            'external_dep': [],
            'precondition': c.get('선행조건') or g.get('선행조건') or None,
            '확인필요': g.get('확인필요') or None,
        },
        'raw': {'ia-feature-canon': c, 'fit-gap-matrix': g},
    })
for r in rows:
    r['근거'] = [x for x in r['근거'] if x]

# ── 원천 3: 260616 IA 엑셀 02_IA마스터 144건 ──────────────────────
import openpyxl
ws = openpyxl.load_workbook(rp(P_XLSX), data_only=True)['02_IA마스터']
hdr = [c.value for c in ws[1]]
n_ia = 0
for ridx, vals in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    rec = {h: v for h, v in zip(hdr, vals)}
    if rec.get('No') is None:
        continue
    n_ia += 1
    no = int(rec['No'])
    rows.append({
        'legacy_id': f'IA-{no:03d}',
        '원천': 'ia',
        '원본ID': f'260616_IA#02_IA마스터!No={no}',
        '기능서술': ' / '.join(str(x) for x in (rec.get('시스템'), rec.get('영역'), rec.get('기능')) if x),
        '원판정': rec.get('진행상태') or '',
        '원판정상세': {
            '진행상태': rec.get('진행상태'),
            '우선순위': rec.get('우선순위'),
            'Phase': rec.get('Phase'),
            '담당': rec.get('담당'),
            '개발규모': rec.get('개발규모'),
        },
        '근거': [f'{P_XLSX}#02_IA마스터!A{ridx}:K{ridx}'],
        '돈여부': None,
        '주문여부': None,
        '의존': {
            'blocked_by': [],
            'external_dep': [],
            'precondition': rec.get('확인·결정 필요사항 (선행조건)'),
            '비고': rec.get('비고'),
        },
        'raw': rec,
    })

json.dump(rows, open(os.path.join(OUT, 'legacy-normalized.json'), 'w', encoding='utf-8'),
          ensure_ascii=False, indent=1)

with open(os.path.join(OUT, 'id-preservation-map.csv'), 'w', newline='', encoding='utf-8') as f:
    w = csv.writer(f)
    w.writerow(['legacy_id', '원천', '원본ID', 'id_unchanged', '기능서술'])
    for r in rows:
        w.writerow([r['legacy_id'], r['원천'], r['원본ID'],
                    'Y' if r['legacy_id'] == r['원본ID'] else 'N',
                    r['기능서술']])

print(f'runway={len(ledger)} scope={len(canon)} ia={n_ia} total={len(rows)}')
