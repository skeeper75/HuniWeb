"""권위 상품마스터에서 키워드가 든 행을 그대로 찾아 내리는 결정론 검색기. 해석하지 않는다."""
import sys

import openpyxl

XLSX = '/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_상품마스터_260822_1.xlsx'


def main():
    if len(sys.argv) == 1:
        wb = openpyxl.load_workbook(XLSX, data_only=True, read_only=True)
        for name in wb.sheetnames:
            print(name)
        return
    needle = sys.argv[1]
    only = sys.argv[2] if len(sys.argv) > 2 else None
    wb = openpyxl.load_workbook(XLSX, data_only=True, read_only=True)
    for name in wb.sheetnames:
        if only and name != only:
            continue
        ws = wb[name]
        for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = ['' if v is None else str(v).replace('\t', ' ').replace('\n', '⏎')
                     for v in row]
            if any(needle in c for c in cells):
                while cells and cells[-1] == '':
                    cells.pop()
                print(f'[{name}] r{r}\t' + '\t'.join(cells))


if __name__ == '__main__':
    main()
