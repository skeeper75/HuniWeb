"""이름표 축 격자 생성기 v3 — 모양 레시피 + 리드 확정 매핑표를 쓴다.

이름표를 코드로 옮기는 순서(리드 4단 사다리):
  ⓪ 리드가 확정한 매핑표(`m3/label-code-map.csv`) — 있으면 그것을 쓴다
  ① 그 그릇을 갈음할 옛 그릇이 실제로 쓰는 코드 중 이름 정규화 유일 일치
  ② 전역(t_siz_sizes · t_mat_materials · t_prd_product_options) 유일 일치
  ③ 없음 → 행을 만들지 않고 「묻는 건」에 적는다 (추측하지 않는다)

동가 묶음(「무광(화이트/블랙)」)은 매핑표에서 `|` 로 갈라 적고 항목별 행으로 편다.
매핑표의 코드 칸이 비어 있으면 그 행은 의도적으로 만들지 않는다(사유를 함께 적는다).
"""
import csv
import re
import sys
import openpyxl

WT = '/Users/innojini/Dev/HuniWeb/.claude/worktrees/t34'
BASE = f'{WT}/_workspace/price-setup'
sys.path.insert(0, f'{BASE}/scripts')
import db                                                        # noqa: E402

XLSX = '/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_인쇄상품_가격표_260902_2.xlsx'
PRICE_TITLE = {'PRICE_TYPE.01': '단가', 'PRICE_TYPE.02': '합가', 'PRICE_TYPE.03': '고정'}
LABEL = {'siz_cd': '사이즈', 'mat_cd': '자재', 'opt_cd': '옵션코드', 'proc_cd': '공정',
         'bdl_qty': '묶음수', 'min_qty': '수량(이상)', 'siz_width': '사이즈가로(이하)',
         'siz_height': '사이즈세로(이하)', 'plt_siz_cd': '판형사이즈',
         'print_opt_cd': '인쇄옵션'}
TBL = {'siz_cd': ('t_siz_sizes', 'siz_nm', 'siz_cd'),
       'mat_cd': ('t_mat_materials', 'mat_nm', 'mat_cd'),
       'opt_cd': ('t_prd_product_options', 'opt_nm', 'opt_cd')}


def norm(v):
    return '' if v is None else str(v).strip()


def key(s):
    """이름 정규화 — 공백·구분자·mm 을 없애고 소문자로."""
    return re.sub(r'[\s\*×xX/()]|mm', '', norm(s)).lower()


def keys(s):
    """맞춰 볼 후보 열쇠들. 「A2」와 「A2 (420X594mm)」, 「우드행거+면끈」과
    「우드행거+면끈 포함」처럼 한쪽이 다른 쪽의 머리인 경우를 잡는다."""
    base = norm(s)
    out = [key(base)]
    head = base.split('(')[0].strip()
    if head and key(head) != out[0]:
        out.append(key(head))
    return [k for k in out if k]


def match(src, label):
    """정규화 열쇠 또는 머리-접두로 유일하게 걸리는 코드."""
    for k in keys(label):
        hit = src.get(k)
        if hit and len(hit) == 1:
            return next(iter(hit))
    for k in keys(label):
        cand = {c for n, codes in src.items() for c in codes
                if n.startswith(k) or k.startswith(n)}
        if len(cand) == 1:
            return next(iter(cand))
    return None


def qty_of(label):
    m = re.match(r'\s*([\d,]+)', norm(label))
    return m.group(1).replace(',', '') if m else None


class Resolver:
    def __init__(self, code, old_comp):
        self.code = code
        self.manual = {}
        self.notes = {}
        for r in csv.DictReader(open(f'{BASE}/m3/label-code-map.csv')):
            if r['new_code'] == code:
                k = (r['axis'], key(r['label']))
                self.manual[k] = [c for c in r['codes'].split('|') if c]
                self.notes[k] = r['note']
        self.local, self.glob = {}, {}
        for axis, (t, nm, cd) in TBL.items():
            loc = {}
            for line in db.q(f"SELECT DISTINCT b.{nm}, a.{axis} FROM t_prc_component_prices a "
                             f"JOIN {t} b ON b.{cd}=a.{axis} WHERE a.comp_cd='{old_comp}'",
                             tuples=True).splitlines():
                if line.strip():
                    n, c = line.split('\t')
                    loc.setdefault(key(n), set()).add(c)
            self.local[axis] = loc
        self.unresolved = []

    def _glob(self, axis):
        if axis not in self.glob:
            t, nm, cd = TBL[axis]
            g = {}
            for line in db.q(f"SELECT {nm}, {cd} FROM {t}", tuples=True).splitlines():
                if line.strip():
                    n, c = line.split('\t')
                    g.setdefault(key(n), set()).add(c)
            self.glob[axis] = g
        return self.glob[axis]

    def resolve(self, axis, label):
        """→ (코드 목록, 생략여부). 코드 목록이 비고 생략이 아니면 미해결이다."""
        k = (axis, key(label))
        if k in self.manual:
            codes = self.manual[k]
            return codes, (not codes)                    # 빈 칸 = 의도된 생략
        for src in (self.local.get(axis, {}), self._glob(axis)):
            cd = match(src, label)
            if cd:
                return [cd], False
        self.unresolved.append((axis, norm(label)))
        return [], False


def cells_colrow(ws, r0, r1, c0, c1, res, col_axis, row_axis):
    out = []
    cols = [(c, norm(ws.cell(r0, c).value)) for c in range(c0 + 1, c1 + 1)]
    cols = [(c, v) for c, v in cols if v]
    for r in range(r0 + 1, r1 + 1):
        rlab = norm(ws.cell(r, c0).value)
        if not rlab:
            continue
        for c, clab in cols:
            v = ws.cell(r, c).value
            if not isinstance(v, (int, float)) or isinstance(v, bool):
                continue
            out.append((clab, rlab, v))
    return out


def cells_vlist(ws, r0, r1, c0, c1):
    out = []
    for r in range(r0 + 1, r1 + 1):
        lab = norm(ws.cell(r, c0).value)
        if not lab:
            continue
        for c in range(c0 + 1, c1 + 1):
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                out.append((lab, '1', v))
                break
    return out


def cells_acryl3col(ws, r0, r1, c0, c1):
    """상품 / 옵션명 / 값 — 상품 칸은 병합돼 비어 있으니 마지막 값을 이어서 쓴다.

    상품 칸이 필요한 이유: 「일자핀」·「2구자석」·「투명」은 라이브에 상품별로 각각
    있어 이름만으로는 코드를 가릴 수 없다(아크릴명찰의 일자핀 = OPV_001025).
    """
    out, prod = [], ''
    for r in range(r0, r1 + 1):
        p = norm(ws.cell(r, c0).value)
        if p:
            prod = p
        lab = norm(ws.cell(r, c0 + 1).value)
        v = ws.cell(r, c0 + 2).value
        if lab and isinstance(v, (int, float)) and not isinstance(v, bool):
            out.append((lab, '1', v, prod))
    return out


def product_opts(prod_label):
    """상품 이름 → 그 상품에 등록된 {정규화이름: {코드}}. 없으면 빈 표."""
    pk = key(prod_label)
    rows = db.q("SELECT prd_cd, prd_nm FROM t_prd_products WHERE use_yn='Y'", tuples=True)
    prd = None
    for line in rows.splitlines():
        if line.strip():
            cd, nm = line.split('\t')
            if key(nm) == pk or key(nm).startswith(pk) or pk.startswith(key(nm)):
                prd = cd
                break
    if not prd:
        return {}
    out = {}
    for line in db.q(f"SELECT opt_nm, opt_cd FROM t_prd_product_options "
                     f"WHERE prd_cd='{prd}' AND coalesce(del_yn,'N')='N'",
                     tuples=True).splitlines():
        if line.strip():
            nm, cd = line.split('\t')
            out.setdefault(key(nm), set()).add(cd)
    return out


def main():
    code = sys.argv[1]
    spec = {r['new_code']: r for r in csv.DictReader(open(f'{BASE}/m2/registration-spec-43.csv'))}
    grids = {g['code']: g for g in csv.DictReader(open(f'{BASE}/m1/authority-grids-trackA.csv'))}
    mapping = {m['new_code']: m for m in csv.DictReader(open(f'{BASE}/m1/mapping-43-trackA.csv'))}
    recipes = {r['new_code']: r for r in csv.DictReader(open(f'{BASE}/m3/shape-recipes.csv'))}
    if code not in recipes:
        print(f'{code}: 레시피 없음 — 건너뜀')
        return 2
    rec = recipes[code]
    g = grids.get(code)
    if not g:
        print(f'{code}: 권위 격자 없음 — 건너뜀')
        return 2
    dims = [d for d in spec[code]['use_dims_to_pick'].split(',') if d]
    sp_prc = spec[code]['prc_typ_cd']
    old = [c.strip() for c in mapping[code]['live_comp'].split(' + ')][0]
    res = Resolver(code, old)

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb[g['sheet']]
    r0, r1 = int(g['grid_head_row']), int(g['grid_last_row'])
    c0, c1 = int(g['grid_col_from']), int(g['grid_col_to'])
    col_axis, row_axis = rec['col_axis'], rec['row_axis']

    if rec['shape'] == 'colrow':
        raw = cells_colrow(ws, r0, r1, c0, c1, res, col_axis, row_axis)
    elif rec['shape'] == 'vlist':
        raw = cells_vlist(ws, r0, r1, c0, c1)
        row_axis = ''
    elif rec['shape'] == 'single':
        raw = []
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                v = ws.cell(r, c).value
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    raw.append(('', '1', v))
        col_axis, row_axis = '', 'min_qty'
    elif rec['shape'] == 'acryl3col':
        raw = cells_acryl3col(ws, r0, r1, c0, c1)
        row_axis = ''
    else:
        print(f'{code}: 모양 {rec["shape"]} 미구현')
        return 2

    rows, skipped = [], []
    pcache = {}
    for item in raw:
        if len(item) == 4:
            clab, rlab, v, prod = item
            # 상품별 옵션표를 먼저 본다(같은 이름이 상품마다 다른 코드를 갖는다).
            if prod not in pcache:
                pcache[prod] = product_opts(prod)
            manual = res.manual.get(('opt_cd', key(f'{prod} / {clab}')))
            if manual is not None:
                if not manual:
                    skipped.append(('opt_cd', f'{prod} / {clab}',
                                    res.notes.get(('opt_cd', key(f'{prod} / {clab}')), '')))
                    continue
                cd = manual[0]
            else:
                cd = match(pcache[prod], clab) if pcache[prod] else None
                if not cd:
                    codes, skip = res.resolve('opt_cd', clab)
                    cd = codes[0] if codes else None
                    if cd:
                        res.unresolved = [u for u in res.unresolved
                                          if u != ('opt_cd', norm(clab))]
            if cd:
                rows.append([''] + [{'opt_cd': cd, 'min_qty': '1'}.get(d, '') for d in dims]
                            + [f'{v:g}', f"권위 {g['sheet']} r{r0}~r{r1} · {prod}"])
                continue
            res.unresolved.append(('opt_cd', f'{prod} / {clab}'))
            continue
        clab, rlab, v = item
        parts = []           # [(axis, [codes])]
        ok = True
        for axis, lab in ((col_axis, clab), (row_axis, rlab)):
            if not axis:
                continue
            if axis == 'min_qty':
                q = qty_of(lab)
                if q is None:
                    ok = False
                    break
                parts.append((axis, [q]))
                continue
            codes, skip = res.resolve(axis, lab)
            if skip:
                skipped.append((axis, lab, res.notes.get((axis, key(lab)), '')))
                ok = False
                break
            if not codes:
                ok = False
                break
            parts.append((axis, codes))
        if not ok:
            continue
        # 동가 묶음 전개 — 코드가 여럿이면 항목별 행으로 편다
        combos = [{}]
        for axis, codes in parts:
            combos = [dict(c, **{axis: cd}) for c in combos for cd in codes]
        for cell in combos:
            cell.setdefault('min_qty', '1')
            rows.append([''] + [cell.get(d, '') for d in dims] + [f'{v:g}',
                        f"권위 {g['sheet']} r{r0}~r{r1}"])

    uniq = sorted(set(res.unresolved))
    if uniq:
        print(f'{code}: ★ 못 옮긴 이름표 {len(uniq)}종 — 표를 만들지 않는다')
        for a, l in uniq:
            print(f'    {LABEL[a]} 「{l}」')
        return 1
    out = f'{BASE}/grid/{code}.csv'
    with open(out, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['적용일'] + [LABEL[d] for d in dims]
                    + [PRICE_TITLE.get(sp_prc, '단가'), '비고'])
        w.writerows(rows)
    sk = f' · 의도적 생략 {len(set(skipped))}종' if skipped else ''
    print(f'{code}: {len(rows)}행 (권위 {g["numeric_cells"]}칸){sk} -> {out}')
    for a, l, n in sorted(set(skipped)):
        print(f'    생략: {LABEL[a]} 「{l}」 — {n}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
