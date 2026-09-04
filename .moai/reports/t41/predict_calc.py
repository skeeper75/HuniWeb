"""t41 — 5케이스의 예측값을 권위 시트에서 결정론으로 유도하고 실측과 대조한다.

예측 = ⌈수량 ÷ 판걸이수⌉ 판 × (그 판수 구간의 시트 단가)
  판걸이수: 시트 「판걸이수」 r77~r82 (반칼스티커 · 317x440)
  단가    : 시트 「스티커」 반칼 블록 r7~r42, 사이즈별 열 묶음
"""
import csv
import math
import os

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = '/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_인쇄상품_가격표_260903.xlsx'

# 사이즈 → (판걸이수, 시트 열)  · 열은 묶음 첫 열(유포/비코팅/미색)
SIZE = {
    'SIZ_000057': (8, 14),    # A6
    'SIZ_000426': (4, 2),     # A5
    'SIZ_000258': (2, 5),     # A4
}
MATCOL = {'MAT_000584': 0, 'MAT_000609': 0, 'MAT_000611': 0,   # 유포/비코팅/미색
          'MAT_000585': 1, 'MAT_000586': 1,
          'MAT_000371': 2, 'MAT_000372': 2, 'MAT_000590': 2}


def main():
    ws = openpyxl.load_workbook(XLSX, data_only=True)['스티커']
    bands, price = [], {}
    for r in range(7, 43):
        q = int(float(ws.cell(r, 1).value))
        bands.append(q)
        for c in range(2, 20):
            v = ws.cell(r, c).value
            if v is not None:
                price[(c, q)] = float(v)

    def unit(col, pansu):
        """판수 이하 최대 구간의 단가(엔진 min_qty '이상' 규칙과 동형)."""
        hit = [b for b in bands if b <= pansu]
        if not hit:
            return None
        return price.get((col, max(hit)))

    rows = list(csv.DictReader(open(f'{HERE}/sim-after.csv')))
    before = {r['case']: r for r in csv.DictReader(open(f'{HERE}/sim-before.csv'))}
    out, bad = [], 0
    for r in rows:
        siz, mat, qty = r['siz_cd'], r['mat_cd'], int(r['qty'])
        pansu_cnt, col0 = SIZE[siz]
        col = col0 + MATCOL[mat]
        pansu = math.ceil(qty / pansu_cnt)
        u = unit(col, pansu)
        expect = None if u is None else int(round(pansu * u))
        actual = int(float(r['final_price'] or 0))
        ok = (expect == actual)
        bad += (not ok)
        out.append(dict(case=r['case'], siz_cd=siz, mat_cd=mat, qty=qty,
                        pansu_per_plate=pansu_cnt, pansu=pansu,
                        sheet_col=f'c{col}', sheet_unit=('' if u is None else int(u)),
                        expect=('' if expect is None else expect),
                        actual=actual,
                        before=int(float(before[r['case']]['final_price'] or 0)),
                        match='일치' if ok else '★불일치'))
    path = f'{HERE}/predict_calc.csv'
    with open(path, 'w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(out[0].keys()))
        w.writeheader()
        w.writerows(out)
    hdr = f"{'케이스':16} {'판걸이':>4} {'판수':>5} {'열':>4} {'단가':>7} {'예측':>9} {'실측':>9} {'적용전':>9}  판정"
    print(hdr)
    for o in out:
        print(f"{o['case']:16} {o['pansu_per_plate']:>4} {o['pansu']:>5} {o['sheet_col']:>4} "
              f"{o['sheet_unit']:>7} {o['expect']:>9} {o['actual']:>9} {o['before']:>9}  {o['match']}")
    print(f"\n5케이스 · 불일치 {bad} -> {path}")


if __name__ == '__main__':
    main()
