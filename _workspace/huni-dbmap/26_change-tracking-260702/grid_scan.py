import openpyxl
from openpyxl.utils import get_column_letter

def load_grid(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    grids={}
    for ws in wb.worksheets:
        rows=[]
        for r in ws.iter_rows(values_only=True):
            rows.append(r)
        # trim fully-empty trailing rows
        while rows and all(c is None for c in rows[-1]):
            rows.pop()
        grids[ws.title]=rows
    wb.close()
    return grids

def norm(v):
    if v is None: return ""
    if isinstance(v,float) and v.is_integer(): return str(int(v))
    return str(v).strip()

def nonempty_rows(rows):
    return sum(1 for r in rows if any(c is not None for c in r))

def scan(tag, oldp, newp):
    print("#"*70); print(tag)
    go=load_grid(oldp); gn=load_grid(newp)
    for sh in go:
        ro=go[sh]; rn=gn.get(sh,[])
        no=nonempty_rows(ro); nn=nonempty_rows(rn)
        maxr=max(len(ro),len(rn))
        maxc=max((max((len(r) for r in ro),default=0)),(max((len(r) for r in rn),default=0)))
        # positional diff
        changed=[]
        for i in range(maxr):
            r_o = ro[i] if i<len(ro) else ()
            r_n = rn[i] if i<len(rn) else ()
            for j in range(maxc):
                vo = norm(r_o[j]) if j<len(r_o) else ""
                vn = norm(r_n[j]) if j<len(r_n) else ""
                if vo!=vn:
                    changed.append((i+1,j+1,vo,vn))
        flag = "  <== CHANGED" if (changed or len(ro)!=len(rn)) else ""
        print(f"[{sh}] rows old={len(ro)}(ne{no}) new={len(rn)}(ne{nn})  pos-changed-cells={len(changed)}{flag}")
        if changed:
            for (i,j,vo,vn) in changed[:40]:
                print(f"    {get_column_letter(j)}{i}: {vo!r} -> {vn!r}")
            if len(changed)>40: print(f"    ... +{len(changed)-40} more")

scan("MASTER 260610->260702",
     "/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_상품마스터_260610.xlsx",
     "/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_상품마스터_260702.xlsx")
scan("PRICE 260527->260702",
     "/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_인쇄상품_가격표_260527.xlsx",
     "/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_인쇄상품_가격표_260702.xlsx")
