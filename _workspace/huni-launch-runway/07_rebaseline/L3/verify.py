#!/usr/bin/env python3
"""L3 완료조건 3개 기계 검증. 실패하면 exit 1."""
import csv, json, os, sys
ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '../../../..'))
L3 = os.path.join(ROOT, '_workspace/huni-launch-runway/07_rebaseline/L3')
rows = json.load(open(os.path.join(L3, 'legacy-normalized.json'), encoding='utf-8'))
ledger = json.load(open(os.path.join(ROOT, '_workspace/huni-launch-runway/01_scope/ledger.json'), encoding='utf-8'))
canon = list(csv.DictReader(open(os.path.join(ROOT, '_workspace/huni-launch-scope/01_foundation/ia-feature-canon.csv'), encoding='utf-8')))
fitgap = list(csv.DictReader(open(os.path.join(ROOT, '_workspace/huni-launch-scope/02_gap/fit-gap-matrix.csv'), encoding='utf-8')))
import openpyxl
ws = openpyxl.load_workbook(os.path.join(ROOT, 'docs/huni/후니프린팅_통합IA_일정_역할분담_260616.xlsx'), data_only=True)['02_IA마스터']
hdr = [c.value for c in ws[1]]
ia = [dict(zip(hdr, v)) for v in ws.iter_rows(min_row=2, values_only=True) if v[0] is not None]

fail = []
def chk(name, ok, detail=''):
    print(('PASS' if ok else 'FAIL'), name, detail)
    if not ok: fail.append(name)

by = lambda s: [r for r in rows if r['원천'] == s]
# ① 건수 + 유실 0
chk('①-1 총건수 716', len(rows) == 716, f'실제 {len(rows)}')
chk('①-2 원천별 410/162/144',
    (len(by('runway')), len(by('scope')), len(by('ia'))) == (410, 162, 144),
    f"{len(by('runway'))}/{len(by('scope'))}/{len(by('ia'))}")
# 유실 0 = 원본 레코드가 raw 에 필드 단위로 그대로 있는가
lost = []
for src, r in zip(ledger, by('runway')):
    if r['raw'] != src: lost.append(('runway', src['id']))
for i, r in enumerate(by('scope')):
    if r['raw']['ia-feature-canon'] != canon[i] or r['raw']['fit-gap-matrix'] != fitgap[i]:
        lost.append(('scope', canon[i]['no']))
for i, r in enumerate(by('ia')):
    if r['raw'] != ia[i]: lost.append(('ia', ia[i]['No']))
chk('①-3 원본 레코드 필드 유실 0', not lost, f'유실 {len(lost)}건 {lost[:5]}')

# ② 원본ID → legacy_id 역추적 100%
m = {r['원본ID']: r['legacy_id'] for r in rows}
chk('②-1 원본ID 유일', len(m) == len(rows), f'{len(m)}/{len(rows)}')
chk('②-2 legacy_id 유일', len({r['legacy_id'] for r in rows}) == len(rows))
mapcsv = list(csv.DictReader(open(os.path.join(L3, 'id-preservation-map.csv'), encoding='utf-8')))
chk('②-3 보존맵 행수 일치', len(mapcsv) == len(rows), f'{len(mapcsv)}')
chk('②-4 보존맵 ↔ json 왕복 일치',
    {(x['원본ID'], x['legacy_id']) for x in mapcsv} == {(k, v) for k, v in m.items()})
chk('②-5 빈 legacy_id/원본ID 0', all(r['legacy_id'] and r['원본ID'] for r in rows))

# ③ runway 항목 ID 무변경
src_ids = [r['id'] for r in ledger]
out_ids = [r['legacy_id'] for r in by('runway')]
chk('③-1 runway ID 집합 동일', set(src_ids) == set(out_ids),
    f'누락 {sorted(set(src_ids)-set(out_ids))[:5]} 신규 {sorted(set(out_ids)-set(src_ids))[:5]}')
chk('③-2 runway ID 순서까지 동일', src_ids == out_ids)
chk('③-3 runway legacy_id == 원본ID (재발급 0)',
    all(r['legacy_id'] == r['원본ID'] for r in by('runway')))
chk('③-4 runway ID 와 신규 prefix 충돌 0',
    not ({r['legacy_id'] for r in by('scope')} | {r['legacy_id'] for r in by('ia')}) & set(src_ids))

# 판정·병합 금지 확인: 원판정이 원천 값 그대로인가
bad = [r['legacy_id'] for r, s in zip(by('runway'), ledger) if r['원판정'] != s.get('status_now','')]
bad += [r['legacy_id'] for r, s in zip(by('scope'), fitgap) if r['원판정'] != s['판정']]
bad += [r['legacy_id'] for r, s in zip(by('ia'), ia) if r['원판정'] != (s.get('진행상태') or '')]
chk('④ 원판정 무변경(재판정 0)', not bad, f'{bad[:5]}')
chk('④-2 병합 0 (1행=1원본)', len(rows) == len(ledger)+len(canon)+len(ia))
# 돈/주문 플래그 보존
chk('⑤ runway 돈/주문 플래그 보존',
    all(r['돈여부'] == s.get('money_critical') and r['주문여부'] == s.get('order_critical')
        for r, s in zip(by('runway'), ledger)),
    f"돈 {sum(1 for r in by('runway') if r['돈여부'])} 주문 {sum(1 for r in by('runway') if r['주문여부'])}")
print('\nRESULT:', 'ALL PASS' if not fail else f'{len(fail)} FAIL {fail}')
sys.exit(1 if fail else 0)
