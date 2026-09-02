"""반칼 붙여넣기표 — 소재 묶음을 자재로 편다(게이트 통과 뒤에만 쓴다).

권위 격자: 규격 6 × 소재묶음 3 × 수량 36. 소재 묶음은 동가라 자재별 행으로 전개한다.
자재 대응은 리드 확정이고, 각 자재의 옛 그릇 가격 벡터가 그 묶음 열과 전건 일치함을
`kisscut_gate.py` 가 먼저 확인한다(1,728칸 · 어긋남 0).
"""
import csv
import openpyxl

WT = '/Users/innojini/Dev/HuniWeb/.claude/worktrees/t34'
BASE = f'{WT}/_workspace/price-setup'
XLSX = '/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_인쇄상품_가격표_260902_2.xlsx'

GROUPS = {'유포/비코팅/미색': ['MAT_000584', 'MAT_000611', 'MAT_000609'],
          '무광코팅/유광코팅': ['MAT_000585', 'MAT_000586'],
          '투명/홀로그램': ['MAT_000371', 'MAT_000372', 'MAT_000590']}
SPEC_SIZE = {'A5(4판) / 124 x186 mm (4판)': ['SIZ_000007', 'SIZ_000059'],
             'A4(2판)': ['SIZ_000520'],
             'A3(1판)': [],                    # 라이브 반칼에 없음 — 생략·묻는 건
             '90*190(6판)': ['SIZ_000060'],
             'A6(8판)/100*148(8판)': ['SIZ_000518'],
             '90*110(12판)': ['SIZ_000519']}


def norm(v):
    return '' if v is None else str(v).strip()


def main():
    ws = openpyxl.load_workbook(XLSX, data_only=True)['스티커']
    spec_of, grp_of, cur = {}, {}, ''
    for c in range(2, 20):
        v = norm(ws.cell(4, c).value)
        if v:
            cur = v
        spec_of[c], grp_of[c] = cur, norm(ws.cell(5, c).value)
    qtys = [(r, int(ws.cell(r, 1).value)) for r in range(7, 43)
            if isinstance(ws.cell(r, 1).value, (int, float))]

    rows, omitted = [], []
    for c in range(2, 20):
        sizes, mats = SPEC_SIZE.get(spec_of[c], []), GROUPS.get(grp_of[c], [])
        if not sizes:
            omitted.append((spec_of[c], grp_of[c]))
            continue
        for r, q in qtys:
            v = ws.cell(r, c).value
            if not isinstance(v, (int, float)) or isinstance(v, bool):
                continue
            for s in sizes:
                for m in mats:
                    rows.append(['', s, m, str(q), f'{v:g}',
                                 f'권위 스티커 r{r} c{c} · 묶음 {grp_of[c]}'])

    out = f'{BASE}/grid/STK_KISSCUT_PRINT.csv'
    with open(out, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['적용일', '사이즈', '자재', '수량(이상)', '단가', '비고'])
        w.writerows(rows)
    print(f'STK_KISSCUT_PRINT: {len(rows):,}행 -> {out}')
    for s, g in sorted(set(omitted)):
        print(f'  생략: 규격 「{s}」 묶음 「{g}」 — 라이브 반칼에 대응 규격 없음(묻는 건)')


if __name__ == '__main__':
    main()
