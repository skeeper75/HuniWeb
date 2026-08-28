"""게시 상품 × 권위 엑셀 시트 커버리지 스캔 — 결정론, LLM 전사 없음.

[역할 경계] 이것은 **커버리지 스캐너**이지 값 비교기가 아니다.
값 대조(축 ③)는 여전히 parse_excel.py → diff.py 단일 경로가 담당한다(REQ-PC-021).
여기서 판정하는 것은 「어느 마일스톤이 어느 시트를 권위로 삼아야 하는가」뿐이다.

입력 : .moai/state/verify/PRICECONF-M0/milestone-assignment-260828.csv (라이브 실측 분모)
출력 : .moai/state/verify/PRICECONF-M0/sheet-coverage-260828.csv
"""
import collections
import csv
import re
import sys

import openpyxl

ASSIGN = '.moai/state/verify/PRICECONF-M0/milestone-assignment-260828.csv'
OUT = '.moai/state/verify/PRICECONF-M0/sheet-coverage-260828.csv'
BOOKS = [
    ('docs/huni/후니프린팅_인쇄상품_가격표_260822_1.xlsx', '가격표'),
    ('docs/huni/후니프린팅_상품마스터_260822_1.xlsx', '마스터'),
]


def norm(s):
    """공백·괄호·슬래시·가운뎃점을 지운 비교용 키. 표기 흔들림을 흡수한다."""
    return re.sub(r'[\s()（）/·]', '', str(s))


def main():
    rows = list(csv.DictReader(open(ASSIGN, encoding='utf-8'), delimiter='|'))
    prods = [(r['milestone'], r['prd_nm']) for r in rows]

    hits = collections.defaultdict(collections.Counter)  # sheet -> milestone -> n
    found = collections.defaultdict(set)                 # prd_nm -> {sheet}

    for path, tag in BOOKS:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        for sn in wb.sheetnames:
            # read_only 모드에서는 max_row 팽창 함정이 없다(engine-map 도구 함정표)
            blob = {norm(v) for row in wb[sn].iter_rows(values_only=True)
                    for v in row if isinstance(v, str) and v.strip()}
            big = '\n'.join(blob)
            for ms, pn in prods:
                if norm(pn) and norm(pn) in big:
                    hits[f'{tag}:{sn}'][ms] += 1
                    found[pn].add(f'{tag}:{sn}')
        wb.close()

    with open(OUT, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['sheet', 'milestone', 'matched_products'])
        for sh, c in sorted(hits.items(), key=lambda x: -sum(x[1].values())):
            for ms, n in sorted(c.items(), key=lambda x: int(x[0][1:])):
                w.writerow([sh, ms, n])

    print(f"{'시트':<28} {'적중':>4}  마일스톤 분포")
    print('-' * 88)
    for sh, c in sorted(hits.items(), key=lambda x: -sum(x[1].values())):
        dist = ' '.join(f'{m}:{n}' for m, n in sorted(c.items(), key=lambda x: int(x[0][1:])))
        print(f'{sh:<28} {sum(c.values()):>4}  {dist}')

    miss = [(m, p) for m, p in prods if p not in found]
    print(f'\n권위 엑셀 어느 시트에도 이름이 없는 게시 상품: {len(miss)}건 / {len(prods)}')
    for m, p in miss:
        print(f'  {m}  {p}')
    print(f'\n→ {OUT}')
    return 0 if not miss else 0  # 미발견은 보고 대상이지 실패가 아니다(§4.2 미게시 조항)


if __name__ == '__main__':
    sys.exit(main())
