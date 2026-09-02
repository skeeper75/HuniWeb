"""M3 붙여넣기표 생성기 — 권위 격자를 단가표 그리드 행으로 편다 (REQ-007).

숫자를 사람이나 모델이 옮겨 적지 않는다. 권위 xlsx 를 직접 읽어 격자를 펴고,
그리드 컬럼 순서(적용일 · <use_dims 순서> · 단가 · 비고)에 맞춰 CSV 로 낸다.

컬럼 순서 근거: comp_price_grid_page.html:221~231 — 적용일(calendar) 이 맨 앞,
그다음 use_dims 순서대로, fk 축은 dropdown 이며 **저장값은 코드(id)** 다.
"""
import csv
import sys
import openpyxl

WT = '/Users/innojini/Dev/HuniWeb/.claude/worktrees/t34'
BASE = f'{WT}/_workspace/price-setup'
XLSX = '/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_인쇄상품_가격표_260902_2.xlsx'


def mm(v):
    """'20mm' · 20 · '20' → 20.0. 격자 머리·라벨을 숫자로 읽는다."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace('mm', '').replace(' ', '')
    try:
        return float(s)
    except ValueError:
        return None


LABEL = {'siz_cd': '사이즈', 'plt_siz_cd': '판형사이즈', 'print_opt_cd': '인쇄옵션',
         'mat_cd': '자재', 'proc_cd': '공정', 'opt_cd': '옵션코드',
         'coat_side_cnt': '코팅면수', 'spot_side_cnt': '별색면수', 'bdl_qty': '묶음수',
         'page_cnt': '페이지수', 'siz_width': '사이즈가로(이하)',
         'siz_height': '사이즈세로(이하)', 'min_qty': '수량(이상)'}


def wh_grid(ws, g, mat_cd, min_qty, note, dims):
    """가로 × 세로 격자 → 행 목록.

    **머리 행(1행) = 가로, 라벨 열(A열) = 세로.**

    방향의 권위는 **상품뷰어의 직접입력 범위**다 — 라이브 그릇의 모양이 아니다.
    아트프린트포스터(`PRD_000118`)는 가로 200~1200 · 세로 200~3000 이고, 권위 인화지
    격자는 열 600~1200 · 행 600~3000 이다. 열이 가로, 행이 세로다. 일반현수막
    (`PRD_000138`, 가로 500~1750 · 세로 500~5000)도 같은 방향이다.

    한때 라이브 그릇의 축 종수(`siz_width` 13종 ≤3000)를 근거로 반대로 읽었으나,
    그것은 **라이브의 전치 적재라는 기존 결함을 정답으로 삼은 순환논증**이었다.
    실무진 확인도 같다 — 「가로를 가로로 세로를 세로로. 미니파츠 때문에 1.5T
    가격표 가로 가격을 더 채워뒀다」(120 열 = 가로).
    """
    r0, r1 = int(g['grid_head_row']), int(g['grid_last_row'])
    c0, c1 = int(g['grid_col_from']), int(g['grid_col_to'])
    widths = [(c, mm(ws.cell(r0, c).value)) for c in range(c0 + 1, c1 + 1)]
    widths = [(c, w) for c, w in widths if w is not None]
    rows = []
    for r in range(r0 + 1, r1 + 1):
        h = mm(ws.cell(r, c0).value)
        if h is None:
            continue
        for c, w in widths:
            v = ws.cell(r, c).value
            if not isinstance(v, (int, float)) or isinstance(v, bool):
                continue
            cell = {'mat_cd': mat_cd, 'siz_width': f'{w:g}', 'siz_height': f'{h:g}',
                    'min_qty': str(min_qty)}
            rows.append([''] + [cell.get(d, '') for d in dims] + [f'{v:g}', note])
    return rows


def main():
    code = sys.argv[1]
    mat_cd = sys.argv[2] if len(sys.argv) > 2 else ''
    min_qty = sys.argv[3] if len(sys.argv) > 3 else '1'
    grids = {g['code']: g for g in csv.DictReader(open(f'{BASE}/m1/authority-grids-trackA.csv'))}
    g = grids[code]
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb[g['sheet']]
    note = f"권위 {g['sheet']} r{g['grid_head_row']}~r{g['grid_last_row']}"
    # 컬럼은 그 그릇의 use_dims 순서를 그대로 따른다(화면 그리드가 그렇게 만들어진다).
    spec = {r['new_code']: r for r in csv.DictReader(open(f'{BASE}/m2/registration-spec-43.csv'))}
    dims = [d for d in spec[code]['use_dims_to_pick'].split(',') if d]
    rows = wh_grid(ws, g, mat_cd, min_qty, note, dims)

    out = f'{BASE}/grid/{code}.csv'
    with open(out, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['적용일'] + [LABEL[d] for d in dims] + ['단가', '비고'])
        w.writerows(rows)
    prices = [float(r[-2]) for r in rows]
    print(f'{code}: {len(rows)}행 -> {out}')
    print(f'  권위 격자 숫자칸 {g["numeric_cells"]} · 만든 행 {len(rows)} · '
          f'{"일치 ✓" if int(g["numeric_cells"]) == len(rows) else "★불일치"}')
    print(f'  단가 범위 {min(prices):g} ~ {max(prices):g}')


if __name__ == '__main__':
    main()
