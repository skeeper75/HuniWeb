"""t40 게이트 — 사후 셀대조(시트 vs 라이브).

커밋하지 않고 사후 상태를 만든다: 한 psql 세션에서
  BEGIN → apply INSERT → 전 행 CSV 추출 → ROLLBACK
그 추출물을 권위 시트 셀과 전수 대조한다. 라이브에는 아무것도 남지 않는다.

시트 열 묶음 (r4/r5 실측):
  c2~c4  A5(4판)/124x186(4판)   c5~c7  A4(2판)     c8~c10 A3(1판)
  c11~c13 90*190(6판)           c14~c16 A6(8판)/100*148(8판)   c17~c19 90*110(12판)
각 묶음 안: 첫 열 유포/비코팅/미색 · 둘째 무광코팅/유광코팅 · 셋째 투명/홀로그램
"""
import csv
import io
import os
import subprocess
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import dbq

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = '/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_인쇄상품_가격표_260903.xlsx'
COMP = 'STK_KISSCUT_PRINT'

GROUP_MATS = [('유포/비코팅/미색', ['MAT_000584', 'MAT_000609', 'MAT_000611']),
              ('무광코팅/유광코팅', ['MAT_000585', 'MAT_000586']),
              ('투명/홀로그램', ['MAT_000371', 'MAT_000372', 'MAT_000590'])]
# 라이브 사이즈 → 시트 묶음 시작 열
SIZE_COL0 = {
    'SIZ_000426': 2,    # A5
    'SIZ_000059': 2,    # 124x186 — A5 와 같은 묶음
    'SIZ_000258': 5,    # A4
    'SIZ_000060': 11,   # 90x190
    'SIZ_000057': 14,   # A6
    'SIZ_000058': 14,   # 100x140 (시트 '100*148' 오타 · t40 신규)
    'SIZ_000067': 14,   # 140x100 (판걸이 8 동일 · t40 신규)
    'SIZ_000519': 17,   # 90x110
}
ORPHAN = {'SIZ_000518': '100x148 — 시트 100*148 은 100x140 오타로 판정됨(카드 ④ 고아)'}


def sheet_cells():
    ws = openpyxl.load_workbook(XLSX, data_only=True)['스티커']
    cells, qtys = {}, []
    for r in range(7, 43):
        q = int(float(ws.cell(r, 1).value))
        qtys.append(q)
        for c in range(2, 20):
            v = ws.cell(r, c).value
            if v is not None:
                cells[(c, q)] = float(v)
    return cells, qtys


def post_state():
    """BEGIN → INSERT → 추출 → ROLLBACK. 한 세션 안에서 끝난다."""
    insert = open(f'{HERE}/apply.sql').read()
    insert = insert.split('\nBEGIN;\n', 1)[1].split('\n-- 검산:', 1)[0].strip()
    sql = (f"BEGIN;\n{insert}\n"
           f"COPY (SELECT siz_cd, mat_cd, min_qty, unit_price "
           f"FROM t_prc_component_prices WHERE comp_cd='{COMP}' "
           f"ORDER BY siz_cd, mat_cd, min_qty) TO STDOUT WITH CSV;\nROLLBACK;\n")
    path = f'{HERE}/_post_export.sql'
    with open(path, 'w') as f:
        f.write(sql)
    r = subprocess.run(['psql', dbq.URL, '-v', 'ON_ERROR_STOP=1', '-q', '-f', path],
                       capture_output=True, text=True)
    if r.returncode:
        raise SystemExit(r.stderr[:800])
    os.remove(path)
    rows = {}
    for line in csv.reader(io.StringIO(r.stdout)):
        if len(line) == 4:
            rows[(line[0], line[1], int(line[2]))] = float(line[3])
    return rows


def main():
    cells, qtys = sheet_cells()
    live = post_state()
    lines = []
    total_checked = mismatch = 0
    per_size = []
    for siz in sorted({k[0] for k in live}):
        got = {(m, q): v for (s, m, q), v in live.items() if s == siz}
        if siz in ORPHAN:
            per_size.append((siz, len(got), '-', '-', f'고아 — {ORPHAN[siz]}'))
            continue
        c0 = SIZE_COL0[siz]
        bad = []
        checked = 0
        for gi, (_, mats) in enumerate(GROUP_MATS):
            col = c0 + gi
            for mat in mats:
                for q in qtys:
                    k = (mat, q)
                    if k not in got:
                        bad.append((mat, q, '라이브 없음', cells.get((col, q))))
                        continue
                    checked += 1
                    want = cells.get((col, q))
                    if want is None or got[k] != want:
                        bad.append((mat, q, got[k], want))
        extra = [k for k in got if k[0] not in
                 [m for _, ms in GROUP_MATS for m in ms]]
        total_checked += checked
        mismatch += len(bad)
        per_size.append((siz, len(got), checked, len(bad),
                         '일치' if not bad and not extra else f'★ {bad[:2]}'))
    lines.append(f'대조 {total_checked}칸 · 불일치 {mismatch}')
    with open(f'{HERE}/celldiff.txt', 'w') as f:
        f.write(f'{"siz_cd":12} {"행":>5} {"대조칸":>6} {"불일치":>6}  판정\n')
        for s, n, c, b, v in per_size:
            f.write(f'{s:12} {n:>5} {str(c):>6} {str(b):>6}  {v}\n')
        f.write('\n' + '\n'.join(lines) + '\n')
    print(open(f'{HERE}/celldiff.txt').read())


if __name__ == '__main__':
    main()
