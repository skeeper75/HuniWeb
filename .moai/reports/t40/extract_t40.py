"""t40 — 권위 「스티커」 시트에서 누락 3곳의 단가행을 뽑아 적재 CSV 로 만든다.

숫자를 사람이나 모델이 옮겨 적지 않는다. 시트를 직접 읽는다.

출처 열(리드 카드 · 지니 확인):
  r4 c14 「A6(8판)/100*148(8판)」 = 이 한 열 묶음이 세 사이즈의 값 원천이다.
    c14 유포/비코팅/미색 · c15 무광코팅/유광코팅 · c16 투명/홀로그램
  시트의 '100*148' 은 100x140 오타(지니 확인) · 140x100 은 같은 판걸이 8 이라 동일 열 적용.
데이터 행: r7~r42 (36 수량구간) — 라이브 기존 조합의 36 구간과 같은 축.
"""
import csv
import os

import openpyxl

XLSX = '/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_인쇄상품_가격표_260903.xlsx'
OUT = os.path.dirname(os.path.abspath(__file__))
COMP = 'STK_KISSCUT_PRINT'
APPLY_YMD = '2026-09-02'          # 기존 1,944행과 동일(카드 지정)
R0, R1 = 7, 42                     # 데이터 행
QCOL = 1
# 열 → (소재 묶음명, 자재 코드) — 카드 본문 재질 매핑 그대로
COLMAT = {
    14: ('유포/비코팅/미색', ['MAT_000584', 'MAT_000609', 'MAT_000611']),
    15: ('무광코팅/유광코팅', ['MAT_000585', 'MAT_000586']),
    16: ('투명/홀로그램', ['MAT_000371', 'MAT_000372', 'MAT_000590']),
}
# 채울 대상 — (사이즈, 넣을 자재 집합 or None=전 자재, 비고 꼬리표)
TARGETS = [
    ('SIZ_000057', {'MAT_000609', 'MAT_000611'}, ''),          # ① A6 — 나머지 6종 기존
    ('SIZ_000058', None,                                        # ② 100x140
     ' · 시트 100*148 은 100x140 오타(지니 확인 260904)'),
    ('SIZ_000067', None,                                        # ③ 140x100
     ' · 판걸이 8 동일로 A6/100*148 열 적용(실무진 확인 필요)'),
]


def main():
    ws = openpyxl.load_workbook(XLSX, data_only=True)['스티커']
    # 머리 검증 — 열이 밀렸으면 여기서 멈춘다
    head = str(ws.cell(4, 14).value or '').strip()
    assert 'A6(8판)' in head, f'출처 열 머리 불일치: {head!r}'
    for c, (want, _) in COLMAT.items():
        got = str(ws.cell(5, c).value or '').strip()
        assert got == want, f'c{c} 소재 머리 불일치: {got!r} != {want!r}'

    src = {}          # (col, qty) -> (price, sheet_row)
    qtys = []
    for r in range(R0, R1 + 1):
        q = ws.cell(r, QCOL).value
        if q is None:
            raise SystemExit(f'r{r} 수량 칸이 비었다 — 데이터 행 범위 재확인 필요')
        q = int(float(q))
        qtys.append(q)
        for c in COLMAT:
            v = ws.cell(r, c).value
            if v is None:
                raise SystemExit(f'r{r} c{c} 값이 비었다')
            src[(c, q)] = (float(v), r)
    assert len(qtys) == len(set(qtys)) == 36, f'수량구간 {len(qtys)}개 (36 기대)'

    rows = []
    for siz, only, tail in TARGETS:
        for c, (group, mats) in COLMAT.items():
            for mat in mats:
                if only is not None and mat not in only:
                    continue
                for q in qtys:
                    price, srow = src[(c, q)]
                    rows.append({
                        'comp_cd': COMP, 'apply_ymd': APPLY_YMD, 'siz_cd': siz,
                        'mat_cd': mat, 'min_qty': q, 'unit_price': f'{price:.2f}',
                        'note': f'권위 스티커 r{srow} c{c} · 묶음 {group}{tail}',
                        'src_col': f'c{c}', 'src_row': srow,
                    })
    path = f'{OUT}/fill_rows.csv'
    with open(path, 'w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    by_siz = {}
    for r in rows:
        by_siz[r['siz_cd']] = by_siz.get(r['siz_cd'], 0) + 1
    print(f'수량구간 {len(qtys)} ({qtys[0]}~{qtys[-1]})')
    print(f'적재행 {len(rows)} · 사이즈별 {by_siz} -> {path}')


if __name__ == '__main__':
    main()
