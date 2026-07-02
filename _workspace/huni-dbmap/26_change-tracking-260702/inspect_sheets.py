import openpyxl, sys
def info(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out=[]
    for ws in wb.worksheets:
        out.append((ws.title, ws.max_row, ws.max_column))
    wb.close()
    return out
pairs = {
 "MASTER": ("/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_상품마스터_260610.xlsx",
            "/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_상품마스터_260702.xlsx"),
 "PRICE":  ("/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_인쇄상품_가격표_260527.xlsx",
            "/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_인쇄상품_가격표_260702.xlsx"),
}
for tag,(a,b) in pairs.items():
    print("="*60); print(tag, "OLD:", a.split('/')[-1])
    oa=info(a); ob=info(b)
    na=[x[0] for x in oa]; nb=[x[0] for x in ob]
    print("--- OLD sheets ---")
    for t,r,c in oa: print(f"  {t!r:45} rows={r} cols={c}")
    print("--- NEW sheets ---")
    for t,r,c in ob: print(f"  {t!r:45} rows={r} cols={c}")
    print("  only OLD:", set(na)-set(nb))
    print("  only NEW:", set(nb)-set(na))
