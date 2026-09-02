"""권위 가격표 시트를 그대로 TSV 로 내리는 결정론 덤퍼. 값을 해석하지 않는다."""
import sys
import openpyxl

XLSX = '/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_인쇄상품_가격표_260902_2.xlsx'

def main():
    sheet = sys.argv[1]
    r0 = int(sys.argv[2]) if len(sys.argv) > 2 else 1
    r1 = int(sys.argv[3]) if len(sys.argv) > 3 else 0
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb[sheet]
    if not r1:
        r1 = ws.max_row
    print(f'# {sheet}  max_row={ws.max_row}  max_col={ws.max_column}', file=sys.stderr)
    for r in range(r0, min(r1, ws.max_row) + 1):
        cells = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            cells.append('' if v is None else str(v).replace('\t', ' ').replace('\n', '⏎'))
        while cells and cells[-1] == '':
            cells.pop()
        print(f'r{r}\t' + '\t'.join(cells))

if __name__ == '__main__':
    main()
