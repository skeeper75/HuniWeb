import openpyxl, csv, warnings, re
from openpyxl.utils import get_column_letter
warnings.filterwarnings("ignore")

M_OLD="/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_상품마스터_260610.xlsx"
M_NEW="/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_상품마스터_260702.xlsx"
P_OLD="/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_인쇄상품_가격표_260527.xlsx"
P_NEW="/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_인쇄상품_가격표_260702.xlsx"

def norm(v):
    if v is None: return ""
    if isinstance(v,float) and v.is_integer(): return str(int(v))
    return re.sub(r"\s+"," ",str(v)).strip()

def grid(path,sheet):
    wb=openpyxl.load_workbook(path,data_only=True); ws=wb[sheet]
    g=[[norm(c) for c in r] for r in ws.iter_rows(values_only=True)]
    wb.close(); return g

def pos_diff(go,gn):
    maxr=max(len(go),len(gn))
    maxc=max(max((len(r) for r in go),default=0),max((len(r) for r in gn),default=0))
    out=[]
    for i in range(maxr):
        ro=go[i] if i<len(go) else []
        rn=gn[i] if i<len(gn) else []
        for j in range(maxc):
            vo=ro[j] if j<len(ro) else ""
            vn=rn[j] if j<len(rn) else ""
            if vo!=vn: out.append((i+1,j+1,vo,vn))
    return out

def hdr_label(go,j,layers=2):
    # combine header rows 1..layers for col j (1-based)
    parts=[]
    for i in range(layers):
        if i<len(go):
            r=go[i]
            if j-1<len(r) and r[j-1]: parts.append(r[j-1])
    return " / ".join(dict.fromkeys(parts))

def ffill_col(go,col_idx,start_row):
    # forward-fill product name from col_idx (0-based), return dict row->name
    names={}; cur=""
    for i in range(start_row,len(go)):
        r=go[i]
        v=r[col_idx] if col_idx<len(r) else ""
        if v: cur=v
        names[i+1]=cur
    return names

# ---------- MASTER diff ----------
mrows=[]
# 디지털인쇄: prd_nm = col D (index3), data starts row3
sh="디지털인쇄"
go=grid(M_NEW,sh); ga=grid(M_OLD,sh)
names=ffill_col(go,3,2)
for (i,j,vo,vn) in pos_diff(ga,go):
    col=hdr_label(go,j,2)
    prd=names.get(i,"")
    mrows.append([sh,prd,col,get_column_letter(j)+str(i),vo,vn])
# MAP: category taxonomy — no prd_nm; key = cell content context
sh="MAP"
go=grid(M_NEW,sh); ga=grid(M_OLD,sh)
for (i,j,vo,vn) in pos_diff(ga,go):
    mrows.append([sh,"(카테고리 트리)","MAP "+get_column_letter(j),get_column_letter(j)+str(i),vo,vn])

with open("master-diff-260610-260702.csv","w",newline="") as f:
    w=csv.writer(f); w.writerow(["sheet","prd_nm","column","cell_ref","before","after"])
    w.writerows(mrows)
print("MASTER diff rows:",len(mrows))

# ---------- PRICE diff (positional sheets) ----------
prows=[]
for sh,keyfn in [("스티커",None),("명함포토카드",None),("포스터사인",None)]:
    go=grid(P_NEW,sh); ga=grid(P_OLD,sh)
    for (i,j,vo,vn) in pos_diff(ga,go):
        col=hdr_label(go,j,3)
        prows.append([sh,"(좌표)",col,get_column_letter(j)+str(i),vo,vn])

# ---------- 출력소재(IMPORT): key-based on 파일명약어(col E idx4), fallback 종이명(C idx2) ----------
sh="출력소재(IMPORT)"
go=grid(P_OLD,sh); gn=grid(P_NEW,sh)
def keyed(g,kidx,fbidx,start=3):
    d={}; order=[]
    for i in range(start,len(g)):
        r=g[i]
        k=r[kidx] if kidx<len(r) else ""
        if not k: k=(r[fbidx] if fbidx<len(r) else "")
        if not k: continue
        # include row values dict by column letter
        d.setdefault(k,[]).append((i+1,r))
    return d
ko=keyed(go,4,2); kn=keyed(gn,4,2)
added=sorted(set(kn)-set(ko)); removed=sorted(set(ko)-set(kn))
hdr=go[0]
import_rows=[]
for k in added:
    for (rr,r) in kn[k]:
        import_rows.append([sh,k,"(ADDED row)","row "+str(rr),"", " | ".join(f"{get_column_letter(j+1)}={v}" for j,v in enumerate(r) if v)[:300]])
for k in removed:
    for (rr,r) in ko[k]:
        import_rows.append([sh,k,"(REMOVED row)","row "+str(rr)," | ".join(f"{get_column_letter(j+1)}={v}" for j,v in enumerate(r) if v)[:300],""])
for k in sorted(set(ko)&set(kn)):
    # compare first occurrence row cells
    ro=ko[k][0][1]; rn=kn[k][0][1]
    mc=max(len(ro),len(rn))
    for j in range(mc):
        vo=ro[j] if j<len(ro) else ""; vn=rn[j] if j<len(rn) else ""
        if vo!=vn:
            col=hdr[j] if j<len(hdr) and hdr[j] else get_column_letter(j+1)
            import_rows.append([sh,k,col,get_column_letter(j+1),vo,vn])

with open("price-diff-260527-260702.csv","w",newline="") as f:
    w=csv.writer(f); w.writerow(["sheet","key","column","cell_ref","before","after"])
    w.writerows(prows+import_rows)
print("PRICE diff rows (positional):",len(prows)," IMPORT key-based rows:",len(import_rows))
print("IMPORT added keys:",added)
print("IMPORT removed keys:",removed)
