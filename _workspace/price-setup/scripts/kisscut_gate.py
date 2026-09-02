"""반칼 소재 묶음 전개 전 검증 게이트 (리드 지시).

권위 반칼 격자의 소재 열은 동가 묶음이다(「유포/비코팅/미색」 등). 묶음을 자재로 펴려면
**각 자재의 옛 그릇 가격 벡터(규격 × 수량)가 그 묶음 열과 전건 일치**해야 한다.
하나라도 어긋나면 전개하지 않고 blocker 로 올린다.
"""
import csv
import sys
import openpyxl

WT = '/Users/innojini/Dev/HuniWeb/.claude/worktrees/t34'
sys.path.insert(0, f'{WT}/_workspace/price-setup/scripts')
import db                                                        # noqa: E402

XLSX = '/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_인쇄상품_가격표_260902_2.xlsx'
GROUPS = {                                   # 권위 소재 묶음 → 자재 코드 (리드 확정)
    '유포/비코팅/미색': ['MAT_000584', 'MAT_000611', 'MAT_000609'],
    '무광코팅/유광코팅': ['MAT_000585', 'MAT_000586'],
    '투명/홀로그램': ['MAT_000371', 'MAT_000372', 'MAT_000590'],
}


def norm(v):
    return '' if v is None else str(v).strip()


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb['스티커']
    # 머리 두 줄: r4 = 규격(3열마다 병합) · r5 = 소재 묶음
    spec_of, grp_of = {}, {}
    cur = ''
    for c in range(2, 20):
        v = norm(ws.cell(4, c).value)
        if v:
            cur = v
        spec_of[c] = cur
        grp_of[c] = norm(ws.cell(5, c).value)
    qtys = [(r, ws.cell(r, 1).value) for r in range(7, 43)]
    qtys = [(r, int(q)) for r, q in qtys if isinstance(q, (int, float))]
    print(f'권위: 규격 {len(set(spec_of.values()))}종 · 소재묶음 {len(set(grp_of.values()))}종 '
          f'· 수량 {len(qtys)}구간')

    # 라이브 반칼 규격의 자재별 (siz_nm, min_qty) → 단가
    live = {}
    for line in db.q("""
        SELECT s.siz_nm, p.mat_cd, p.min_qty, p.unit_price
        FROM t_prc_component_prices p JOIN t_siz_sizes s ON s.siz_cd = p.siz_cd
        WHERE p.comp_cd = 'COMP_STK_PRINT'
    """, tuples=True).splitlines():
        if line.strip():
            nm, mat, q, price = line.split('\t')
            live[(nm.strip(), mat, int(q))] = float(price)
    sizes = sorted({k[0] for k in live})
    print(f'라이브 반칼 그릇 규격 {len(sizes)}종: {", ".join(sizes[:8])}…\n')

    # 권위 규격 라벨 → 라이브 반칼 규격. 헐거운 이름 매칭은 완칼 규격까지 끌어와
    # 오판을 만든다(「A4(2판)」이 완칼 A4 에 붙어 4,000 vs 5,000 으로 어긋났다).
    # 반칼 규격은 자재 10종 × 수량 36 = 360행인 것들로 실측해 못박는다.
    SPEC_SIZE = {
        'A5(4판) / 124 x186 mm (4판)': [('SIZ_000007', '148x210mm'),
                                        ('SIZ_000059', '124x186mm')],
        'A4(2판)': [('SIZ_000520', 'A4(210x297mm) 반칼')],
        'A3(1판)': [],                       # 라이브 반칼에 A3 규격이 없다 — 묻는 건
        '90*190(6판)': [('SIZ_000060', '90x190mm')],
        'A6(8판)/100*148(8판)': [('SIZ_000518', '100x148')],
        '90*110(12판)': [('SIZ_000519', '90x110')],
    }

    def match_sizes(label):
        return [nm for _, nm in SPEC_SIZE.get(label, [])]

    bad, checked, unmapped = [], 0, []
    for c in range(2, 20):
        grp, slab = grp_of[c], spec_of[c]
        if grp not in GROUPS:
            continue
        tgt = match_sizes(slab)
        if not tgt:
            unmapped.append(slab)
            continue
        for snm in tgt:
            for mat in GROUPS[grp]:
                for r, q in qtys:
                    v = ws.cell(r, c).value
                    if not isinstance(v, (int, float)):
                        continue
                    lv = live.get((snm, mat, q))
                    checked += 1
                    if lv is None:
                        bad.append((snm, mat, q, '라이브 행 없음', float(v)))
                    elif lv != float(v):
                        bad.append((snm, mat, q, lv, float(v)))
    print(f'대조한 칸 {checked:,} · 어긋남 {len(bad):,}')
    if unmapped:
        print('규격 라벨을 라이브에서 못 찾음:', sorted(set(unmapped)))
    from collections import Counter
    if bad:
        c1 = Counter((b[0], b[1]) for b in bad)
        print('\n어긋난 (규격, 자재) 상위:')
        for (s, m), n in c1.most_common(8):
            print(f'  {s:<22} {m}  {n}칸')
        print('\n표본 5:')
        for b in bad[:5]:
            print(f'  {b[0]} {b[1]} 수량{b[2]}: 라이브 {b[3]} vs 권위 {b[4]}')
    print('\n게이트:', 'PASS — 전개 가능' if not bad and not unmapped else '★ FAIL — 전개 금지')
    return 0 if (not bad and not unmapped) else 1


if __name__ == '__main__':
    sys.exit(main())
