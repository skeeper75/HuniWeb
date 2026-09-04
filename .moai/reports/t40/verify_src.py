"""t40 사전 검증 — 뽑아낸 출처 열이 맞는지 '이미 라이브에 있는' 조합으로 대조한다.

A6(SIZ_000057) 는 6자재가 이미 적재돼 있다. 그 6자재 값이 시트 c14/c15/c16 과
일치하면, 같은 열에서 뽑은 누락분(609·611·100x140·140x100)도 옳다는 뜻이다.
수량구간 축이 라이브와 같은지도 함께 본다.
"""
import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import dbq

XLSX = '/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_인쇄상품_가격표_260903.xlsx'
COLMAT = {14: ['MAT_000584'], 15: ['MAT_000585', 'MAT_000586'],
          16: ['MAT_000371', 'MAT_000372', 'MAT_000590']}

ws = openpyxl.load_workbook(XLSX, data_only=True)['스티커']
sheet = {}
qtys = []
for r in range(7, 43):
    q = int(float(ws.cell(r, 1).value))
    qtys.append(q)
    for c in COLMAT:
        sheet[(c, q)] = float(ws.cell(r, c).value)

live = {}
out = dbq.q("SELECT mat_cd, min_qty, unit_price FROM t_prc_component_prices "
            "WHERE comp_cd='STK_KISSCUT_PRINT' AND siz_cd='SIZ_000057'", tuples=True)
for ln in out.splitlines():
    if ln.strip():
        m, q, u = ln.split('\t')
        live[(m.strip(), int(float(q)))] = float(u)

bad, checked = [], 0
for c, mats in COLMAT.items():
    for mat in mats:
        for q in qtys:
            k = (mat, q)
            if k not in live:
                bad.append((mat, q, 'live 없음', sheet[(c, q)]))
                continue
            checked += 1
            if live[k] != sheet[(c, q)]:
                bad.append((mat, q, live[k], sheet[(c, q)]))

lq = sorted({q for _, q in live})
print(f'라이브 A6 수량구간 {len(lq)} · 시트 {len(qtys)} · 동일 {lq == sorted(qtys)}')
print(f'대조 {checked}칸 · 불일치 {len(bad)}')
for b in bad[:8]:
    print('  ★', b)
print('판정:', '출처 열 확정 ✓' if not bad and lq == sorted(qtys) else '★확인필요')
