"""권위 가격표에서 코드 블록을 결정론으로 전수 훑는다 (v2).

v1 이 놓친 것 둘을 고쳤다.
 1) 라벨 어휘가 시트마다 다르다 — 'code'/'name' 뿐 아니라 '구성요소코드'/'구성요소명'
    을 쓰는 블록이 있다(메쉬현수막). 어휘를 넓히지 않으면 블록 하나가 통째로 빠진다.
 2) 한 행에 블록이 옆으로 여러 개 놓인다(일반현수막: 출력가·가공가·부자재 셋).
    행 구간만으로 자르면 앞의 둘이 크기 0 으로 잡힌다. 열 경계로도 자른다.
"""
import sys
import csv
import openpyxl

XLSX = '/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_인쇄상품_가격표_260902_2.xlsx'
SHEETS = ['아크릴', '스티커', '포스터사인']
CODE_LABELS = {'code', '구성요소코드'}
NAME_LABELS = {'name', '구성요소명'}


def norm(v):
    return '' if v is None else str(v).strip()


def scan(ws):
    hits = []
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            if norm(ws.cell(r, c).value).lower() in CODE_LABELS:
                code = norm(ws.cell(r, c + 1).value)
                name = ''
                for dr in (1, 2, 3):
                    if norm(ws.cell(r + dr, c).value).lower() in NAME_LABELS:
                        name = norm(ws.cell(r + dr, c + 1).value)
                        break
                if code:
                    hits.append([r, c, code, name])
    return hits


def extents(ws, hits):
    """각 블록의 (행끝, 열끝) 을 정한다. 같은 행의 다음 블록이 열 경계가 된다."""
    rows = sorted({h[0] for h in hits})
    for h in hits:
        r, c = h[0], h[1]
        nxt_rows = [x for x in rows if x > r]
        row_end = (nxt_rows[0] - 1) if nxt_rows else ws.max_row
        same_row_cols = sorted(x[1] for x in hits if x[0] == r and x[1] > c)
        col_end = (same_row_cols[0] - 1) if same_row_cols else ws.max_column
        yield h, row_end, col_end


def count_numeric(ws, r0, r1, c0, c1):
    n, first, last = 0, None, None
    for r in range(r0, r1 + 1):
        hit = False
        for c in range(c0, c1 + 1):
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                n += 1
                hit = True
        if hit:
            first = r if first is None else first
            last = r
    return n, first, last


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    out = []
    for sheet in SHEETS:
        ws = wb[sheet]
        hits = scan(ws)
        for (r, c, code, name), row_end, col_end in extents(ws, hits):
            n, first, last = count_numeric(ws, r, row_end, c, col_end)
            out.append(dict(sheet=sheet, code_row=r, code_col=c,
                            code_as_written=code, code_upper=code.upper(), name=name,
                            row_from=r, row_to=row_end, col_from=c, col_to=col_end,
                            numeric_first_row=first or '', numeric_last_row=last or '',
                            numeric_cells=n))
        print(f'{sheet}: 코드 블록 {len(hits)}개', file=sys.stderr)

    path = ('/Users/innojini/Dev/HuniWeb/.claude/worktrees/t34/'
            '_workspace/price-setup/m1/authority-blocks-trackA.csv')
    with open(path, 'w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(out[0].keys()))
        w.writeheader()
        w.writerows(out)
    print(f'합계 {len(out)}블록 -> {path}', file=sys.stderr)
    for o in out:
        print(f"{o['sheet']:6s} r{o['code_row']:<4d}c{o['code_col']:<3d} "
              f"{o['code_upper']:<32s} {o['name']:<22s} "
              f"칸 {o['numeric_cells']:>5d}  r{o['numeric_first_row']}~r{o['numeric_last_row']} "
              f"c{o['col_from']}~c{o['col_to']}")


if __name__ == '__main__':
    main()
