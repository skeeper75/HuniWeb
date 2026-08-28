"""포스터사인 시트 → 권위 격자 CSV (결정론 파서, LLM 전사 없음)."""
import csv, json, sys
import openpyxl

SRC = 'docs/huni/후니프린팅_인쇄상품_가격표_260822_1.xlsx'
OUT = '_workspace/postersign-audit/authority.csv'

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb['포스터사인']
MAXR, MAXC = ws.max_row, ws.max_column  # 접근 시 셀이 생성돼 max_row 가 늘어나므로 먼저 고정

def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)

rows_out = []
blocks = []

# 블록 헤더 탐지: A열 문자열 + 다음 행 A열이 'X / Y' 축 라벨
r = 1
while r <= MAXR:
    a = ws.cell(r, 1).value
    nxt = ws.cell(r + 1, 1).value
    if isinstance(a, str) and a.strip() and isinstance(nxt, str) and '/' in nxt:
        name = a.strip()
        axis = nxt.strip()
        hdr = r + 1
        # 열 라벨 수집
        cols = []
        for c in range(2, MAXC + 1):
            v = ws.cell(hdr, c).value
            if v in (None, ''):
                break
            cols.append((c, str(v).strip()))
        # 값 행 수집
        body = []
        rr = hdr + 1
        while rr <= MAXR:
            rl = ws.cell(rr, 1).value
            if rl in (None, ''):
                break
            vals = [(cl, ws.cell(rr, c).value) for c, cl in cols]
            if not any(is_num(v) for _, v in vals):
                break
            body.append((str(rl).strip(), vals))
            rr += 1
        if body:
            blocks.append(dict(name=name, axis=axis, row=r, cols=[c for _, c in cols],
                               nrow=len(body), ncell=sum(1 for _, vs in body for _, v in vs if is_num(v))))
            for rl, vals in body:
                for cl, v in vals:
                    if is_num(v):
                        rows_out.append(dict(block=name, axis=axis, row_label=rl, col_label=cl,
                                             price=int(v), kind='main'))
            r = rr
            continue
    r += 1

# 추가옵션 블록(J~M): J열에 '추가가격' 문구가 있는 행 기준
r = 1
while r <= MAXR:
    j = ws.cell(r, 10).value
    if isinstance(j, str) and '추가가격' in j:
        title = j.strip()
        hdr = r + 1
        cols = []
        for c in range(11, MAXC + 1):
            v = ws.cell(hdr, c).value
            if v in (None, ''):
                break
            cols.append((c, str(v).strip()))
        if not cols:
            # 열 라벨이 없는 단일 열 추가옵션(예: PET배너 거치대) — K열을 무라벨 열로 취급
            cols = [(11, '*단일')]
        rr = hdr + 1
        while rr <= MAXR:
            rl = ws.cell(rr, 10).value
            if rl in (None, ''):
                break
            # 다음 추가옵션 블록의 제목을 데이터로 먹지 않도록 경계에서 멈춘다
            if isinstance(rl, str) and ('추가가격' in rl or '추가옵션' in rl):
                break
            for c, cl in cols:
                v = ws.cell(rr, c).value
                if is_num(v):
                    rows_out.append(dict(block=title, axis='추가옵션', row_label=str(rl).strip(),
                                         col_label=cl, price=int(v), kind='addon'))
            rr += 1
        r = rr
        continue
    r += 1

with open(OUT, 'w', newline='', encoding='utf-8') as f:
    w = csv.DictWriter(f, fieldnames=['block', 'axis', 'kind', 'row_label', 'col_label', 'price'])
    w.writeheader()
    for d in rows_out:
        w.writerow({k: d[k] for k in w.fieldnames})

print(f'본품 블록 {len(blocks)}개')
for b in blocks:
    print(f"  {b['name']:<28} axis={b['axis']:<12} 행{b['nrow']:>3} 열{len(b['cols']):>2} = {b['ncell']:>3}셀")
print(f"\n총 셀 {len(rows_out)} (본품 {sum(1 for d in rows_out if d['kind']=='main')} · 추가옵션 {sum(1 for d in rows_out if d['kind']=='addon')})")
print(f'→ {OUT}')
