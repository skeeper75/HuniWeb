import openpyxl
from openpyxl.utils import get_column_letter
def dump(path, sheet, hdr_rows=(1,2,3,4,5), cols=None):
    wb=openpyxl.load_workbook(path,data_only=True); ws=wb[sheet]
    grid=[r for r in ws.iter_rows(values_only=True)]
    wb.close()
    print(f"=== {sheet} (first rows) ===")
    for i in hdr_rows:
        if i-1<len(grid):
            r=grid[i-1]
            cells=[(get_column_letter(j+1),v) for j,v in enumerate(r) if v is not None]
            print(f" row{i}:", cells[:30])
# MASTER 디지털인쇄 header + rows around 6,28,76
dump("/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_상품마스터_260702.xlsx","디지털인쇄",(1,2,3,4,5,6,28,76))
print()
# PRICE 출력소재(IMPORT) header
dump("/Users/innojini/Dev/HuniWeb/docs/huni/후니프린팅_인쇄상품_가격표_260702.xlsx","출력소재(IMPORT)",(1,2,3))
